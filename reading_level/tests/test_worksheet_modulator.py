"""Worksheet segmentation, source ceiling, DOK bucketing, and modulation."""

import base64
import io
import json
from types import SimpleNamespace

import pytest

import ai
from reading_level import bands as rl_bands
from reading_level import config as rl_config
from reading_level import dok as rl_dok
from reading_level.blocks import segment_worksheet_bytes
from reading_level.scorer import score_text


WORKSHEET_LINES = [
    "Name: ________",
    "Date: ________",
    "Directions: Read the story. Then answer the questions.",
    (
        "Mia has a red dog. The dog is named Sam. Sam likes to run in the park. "
        "Mia throws a ball for him. Sam runs fast and brings it back. They play "
        "until the sun goes down. Then Mia gives Sam some water. Sam drinks it "
        "all. He wags his tail at her. Mia laughs and pats his head. They walk "
        "home together."
    ),
    "1. What is the dog's name?",
    "A. Sam",
    "B. Mia",
    "C. Park",
    "D. Tail",
    "2. Where do Mia and Sam play?",
    "A. The park",
    "B. The store",
    "C. School",
    "D. The lake",
    "3. Why did Sam wag his tail?",
    "4. What can you infer about how Mia feels about Sam? Support your answer with evidence from the text.",
]


def _sample_png() -> bytes:
    import pymupdf as fitz

    doc = fitz.open()
    page = doc.new_page(width=48, height=48)
    page.draw_rect(page.rect, color=(0.15, 0.45, 0.85), fill=(0.15, 0.45, 0.85))
    data = page.get_pixmap().tobytes("png")
    doc.close()
    return data


TINY_PNG = _sample_png()


def _sample_docx() -> bytes:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_paragraph(WORKSHEET_LINES[0])
    document.add_paragraph(WORKSHEET_LINES[1])
    document.add_paragraph(WORKSHEET_LINES[2])
    document.add_paragraph(WORKSHEET_LINES[3])
    document.add_picture(io.BytesIO(TINY_PNG), width=Inches(0.4))
    for line in WORKSHEET_LINES[4:]:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sample_pdf() -> bytes:
    import pymupdf as fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in WORKSHEET_LINES:
        page.insert_text((72, y), line, fontsize=11)
        y += 18
    page.insert_image(fitz.Rect(400, 80, 448, 128), stream=TINY_PNG)
    data = doc.tobytes()
    doc.close()
    return data


def _blocks_from_lines():
    from reading_level.blocks import WorksheetBlock

    return [
        WorksheetBlock("b0", "ignore", WORKSHEET_LINES[0], 0, "high"),
        WorksheetBlock("b1", "instructions", WORKSHEET_LINES[2], 1, "high"),
        WorksheetBlock("b2", "passage", WORKSHEET_LINES[3], 2, "high"),
        WorksheetBlock("b3", "question", WORKSHEET_LINES[4], 3, "high"),
        WorksheetBlock("b4", "answer_choice", WORKSHEET_LINES[5], 4, "high"),
        WorksheetBlock("b5", "answer_choice", WORKSHEET_LINES[6], 5, "high"),
        WorksheetBlock("b6", "answer_choice", WORKSHEET_LINES[7], 6, "high"),
        WorksheetBlock("b7", "answer_choice", WORKSHEET_LINES[8], 7, "high"),
        WorksheetBlock("b8", "question", WORKSHEET_LINES[9], 8, "high"),
        WorksheetBlock("b9", "answer_choice", WORKSHEET_LINES[10], 9, "high"),
        WorksheetBlock("b10", "answer_choice", WORKSHEET_LINES[11], 10, "high"),
        WorksheetBlock("b11", "answer_choice", WORKSHEET_LINES[12], 11, "high"),
        WorksheetBlock("b12", "answer_choice", WORKSHEET_LINES[13], 12, "high"),
        WorksheetBlock("b13", "question", WORKSHEET_LINES[14], 13, "high"),
        WorksheetBlock("b14", "question", WORKSHEET_LINES[15], 14, "high"),
        WorksheetBlock(
            "b15",
            "image",
            "",
            15,
            "high",
            asset_base64=base64.b64encode(TINY_PNG).decode(),
            asset_mime="image/png",
        ),
    ]


def test_docx_segmentation_labels_the_unambiguous_majority():
    blocks = segment_worksheet_bytes(_sample_docx(), "sample.docx")
    types = [block.block_type for block in blocks]
    assert "instructions" in types
    assert "passage" in types
    assert types.count("question") >= 3
    assert types.count("answer_choice") >= 6
    assert "ignore" in types
    assert "image" in types
    ignore_text = " ".join(block.text.lower() for block in blocks if block.block_type == "ignore")
    assert "name" in ignore_text
    assert any(block.asset_base64 for block in blocks if block.block_type == "image")
    for block in blocks:
        if block.block_type == "ignore" and block.confidence == "low":
            assert len(block.text.split()) < 8


def test_pdf_segmentation_flags_low_confidence_instead_of_guessing():
    blocks = segment_worksheet_bytes(_sample_pdf(), "sample.pdf")
    types = [block.block_type for block in blocks]
    assert "question" in types
    assert "answer_choice" in types
    assert any(
        block.block_type == "ignore" and "name" in block.text.lower()
        for block in blocks
    )
    assert any(block.block_type == "image" and block.asset_base64 for block in blocks)
    # PDF is best-effort: numbered stems/choices can still be high, body is low.
    passages = [block for block in blocks if block.block_type == "passage"]
    assert passages
    assert all(block.confidence == "low" for block in passages)


def test_wrapped_pdf_questions_stay_one_tag_even_with_two_marks():
    """Visual wraps and extra `?` must not split a numbered question tag."""
    import pymupdf as fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (72, 72),
        "1. What can you infer about why Leo smiled all the way home even though he didn't score the",
        fontsize=11,
    )
    page.insert_text(
        (72, 90),
        "winning goal? Support your answer with evidence from the text?",
        fontsize=11,
    )
    page.insert_text(
        (72, 120),
        "2. How do the details about Leo playing at the park with backpacks for a goal help you understand",
        fontsize=11,
    )
    page.insert_text(
        (72, 138),
        "what soccer means to him? Support your answer with evidence from the text?",
        fontsize=11,
    )
    data = doc.tobytes()
    doc.close()

    questions = [
        block
        for block in segment_worksheet_bytes(data, "wrap.pdf")
        if block.block_type == "question"
    ]
    assert len(questions) == 2
    assert "score the winning goal" in questions[0].text
    assert "Support your answer with evidence from the text?" in questions[0].text
    assert "understand what soccer means to him" in questions[1].text
    assert questions[0].text.count("?") == 2
    assert not questions[0].text.startswith("winning")


def test_oversized_source_material_is_rejected_not_truncated():
    huge = "word " * ((rl_config.SOURCE_MATERIAL_MAX_TOKENS * rl_config.SOURCE_CHARS_PER_TOKEN) + 20)
    with pytest.raises(ValueError, match="too long"):
        rl_dok.check_source_material(huge, required=False)


def test_test_worksheets_require_source_material():
    with pytest.raises(ValueError, match="source material"):
        rl_dok.check_source_material("", required=True)


def test_dok_bucketing_is_in_the_right_ballpark():
    passage = WORKSHEET_LINES[3]
    from reading_level.blocks import WorksheetBlock

    recall = WorksheetBlock("q1", "question", "What is the dog's name?", 0, "high")
    why = WorksheetBlock("q2", "question", "Why did Sam wag his tail?", 1, "high")
    infer = WorksheetBlock(
        "q3",
        "question",
        "What can you infer about how Mia feels about Sam? Support your answer with evidence from the text.",
        2,
        "high",
    )
    assert rl_dok.diagnose_question(recall, passage).estimated_dok == 1
    assert rl_dok.diagnose_question(why, passage).estimated_dok == 2
    assert rl_dok.diagnose_question(infer, passage).estimated_dok == 3


def test_single_passage_corrector_path_does_not_call_modulate(monkeypatch):
    """The original /correct-reading-level path stays on correct_with_rewrite."""
    called = {"modulate": False}

    def boom(*args, **kwargs):
        called["modulate"] = True
        raise AssertionError("worksheet modulator ran on single-passage path")

    monkeypatch.setattr(ai, "modulate_worksheet", boom)
    outcome = ai.correct_with_rewrite("Mia has a red dog. The dog is named Sam.", 1)
    assert called["modulate"] is False
    assert outcome.text
    assert outcome.correction is not None


def _fake_correct_outcome(text, target, **kwargs):
    score = score_text(text)
    band = rl_bands.target_band(target)
    return SimpleNamespace(
        text=text,
        correction=SimpleNamespace(initial_score=score, target_band=band),
        final_score=score,
        gate_passed=rl_bands.in_band(score, band),
        llm_rewrite_applied=False,
        target_band=band,
    )


def test_modulate_increase_and_decrease_report(monkeypatch, fixture_text):
    """Full in-memory worksheet, both directions. LLM convert is mocked."""
    converted = []

    def fake_convert(stem, target_dok, *args, **kwargs):
        converted.append((stem, target_dok))
        if target_dok >= 3:
            return (
                "What can you infer about Mia and Sam from the story? "
                "Support your answer with evidence from the text?"
            )
        if target_dok == 2:
            return "Why did Sam run back to Mia with the ball?"
        return "What is the dog's name?"

    monkeypatch.setattr(ai, "convert_question_to_dok", fake_convert)
    monkeypatch.setattr(ai, "correct_with_rewrite", _fake_correct_outcome)
    monkeypatch.setattr(
        ai, "correct_text",
        lambda text, target, **kwargs: SimpleNamespace(text=text, gate_passed=True),
    )

    easy = fixture_text("easy_primary")
    hard = fixture_text("business_register")

    def run(passage, grade):
        blocks = _blocks_from_lines()
        blocks[2] = type(blocks[2])(
            blocks[2].block_id, "passage", passage, 2, "high"
        )
        return ai.modulate_worksheet(
            blocks,
            grade,
            source_material="Chapter notes: Mia and her dog Sam play in the park.",
            worksheet_kind="practice",
            allow_rewrite=False,
            title="Sam",
        )

    decrease = run(hard, 3)
    increase = run(easy, 9)

    assert decrease["dok"]["before_counts"]
    assert increase["dok"]["before_counts"]
    assert decrease["pdf_base64"]
    assert increase["pdf_base64"]
    assert "passages" in decrease and "passages" in increase
    ignore_text = " ".join(
        block["text"] for block in decrease["blocks"] if block["block_type"] == "ignore"
    )
    assert "Name:" in ignore_text
    image_blocks = [
        block for block in increase["blocks"]
        if block["block_type"] == "image" and block.get("asset_base64")
    ]
    assert image_blocks
    import pymupdf as fitz

    pdf = fitz.open(stream=base64.b64decode(increase["pdf_base64"]), filetype="pdf")
    embedded = [image for page in pdf for image in page.get_images()]
    pdf.close()
    assert embedded
    assert converted == []


def test_modulate_converts_questions_only_when_asked(monkeypatch):
    from reading_level.blocks import WorksheetBlock

    converted = []
    infer = (
        "What can you infer about Mia? Support your answer with evidence from the text."
    )

    def fake_convert(stem, target_dok, *args, **kwargs):
        converted.append(stem)
        return "What is the dog's name?"

    monkeypatch.setattr(ai, "convert_question_to_dok", fake_convert)
    monkeypatch.setattr(ai, "correct_with_rewrite", _fake_correct_outcome)
    monkeypatch.setattr(
        ai, "correct_text",
        lambda text, target, **kwargs: SimpleNamespace(text=text, gate_passed=True),
    )
    blocks = [
        WorksheetBlock("p", "passage", "Mia has a red dog named Sam.", 0, "high"),
        WorksheetBlock("q1", "question", infer, 1, "high"),
        WorksheetBlock("q2", "question", infer, 2, "high"),
    ]
    defaulted = ai.modulate_worksheet(blocks, 3, allow_rewrite=False, title="Sam")
    assert converted == []
    assert defaulted["dok"]["converted_block_ids"] == []
    assert defaulted["blocks"][1]["text"] == infer
    assert all(fail["check"] != "dok_mix" for fail in defaulted["gate_failures"])

    asked = ai.modulate_worksheet(
        blocks, 3, allow_rewrite=False, convert_questions=True, title="Sam"
    )
    assert converted
    assert asked["dok"]["converted_block_ids"]


def test_legacy_unclassified_maps_to_ignore():
    from reading_level.blocks import WorksheetBlock

    block = WorksheetBlock.from_dict({
        "block_id": "x",
        "block_type": "unclassified",
        "text": "Name: ____",
        "source_position": 0,
        "confidence": "high",
    })
    assert block.block_type == "ignore"
    assert "Name" in block.text


def test_image_block_round_trips_through_from_dict():
    from reading_level.blocks import WorksheetBlock

    payload = base64.b64encode(TINY_PNG).decode()
    block = WorksheetBlock.from_dict({
        "block_id": "img",
        "block_type": "image",
        "text": "",
        "source_position": 3,
        "confidence": "high",
        "asset_base64": payload,
        "asset_mime": "image/png",
    })
    assert block.block_type == "image"
    assert block.asset_base64 == payload
    assert block.asset_mime == "image/png"


def test_generated_sheet_rebuild_keeps_template_html_and_image_slot():
    import file as worksheet_file
    from reading_level.blocks import WorksheetBlock

    image = base64.b64encode(TINY_PNG).decode()
    blocks = [
        WorksheetBlock("b0", "ignore", "Name", 0, "high"),
        WorksheetBlock("b1", "ignore", "Date", 1, "high"),
        WorksheetBlock("b2", "ignore", "The Big Game", 2, "low"),
        WorksheetBlock("b3", "passage", "Leo loved soccer more than anything.", 3, "high"),
        WorksheetBlock("b4", "question", "1. Why did Leo smile all the way home?", 4, "high"),
        WorksheetBlock("b5", "question", "2. What did Leo use as a goal?", 5, "high"),
        WorksheetBlock(
            "b6",
            "image",
            "",
            6,
            "high",
            asset_base64=image,
            asset_mime="image/png",
        ),
    ]
    payload = worksheet_file.blocks_to_worksheet_data("Leveled worksheet", blocks)
    assert payload["Title"] == "The Big Game"
    assert payload["Story"] == "Leo loved soccer more than anything."
    assert payload["Q1"] == "Why did Leo smile all the way home?"
    html = worksheet_file.build_worksheet_html(payload, include_image_slot=True)
    assert "meta-line" in html
    assert "answer-line" in html
    assert "image-slot" in html
    assert "Leveled worksheet" not in html
    assert html.count("class=\"question\"") == 2

    pdf_bytes = worksheet_file.generate_blocks_pdf("Leveled worksheet", blocks)
    import pymupdf as fitz

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    page = doc[0]
    images = [
        block for block in page.get_text("dict")["blocks"] if block.get("type") == 1
    ]
    doc.close()
    assert images
    bbox = images[0]["bbox"]
    assert abs(bbox[0] - 355) < 2
    assert abs(bbox[1] - 125) < 2
    assert abs((bbox[2] - bbox[0]) - 200) < 2
    assert abs((bbox[3] - bbox[1]) - 200) < 2


def test_convert_prompt_reuses_dok_definitions():
    prompt = ai.build_convert_question_prompt(
        "What is the dog's name?", 3, "Mia has a dog named Sam.", "notes"
    )
    assert "Strategic Thinking" in prompt
    assert "What is the dog's name?" in prompt
    assert "Mia has a dog named Sam." in prompt
    assert "notes" in prompt
